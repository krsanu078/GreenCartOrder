import os
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException


EXPECTED_HEADERS = ("Item", "Quantities", "Promo_code", "Country")

def load_test_data(path=None):
    if path is None:
        path = os.getenv("TEST_DATA_FILE")
    if path is None:
        path = Path(__file__).resolve().parent.parent / "GreenCartOrder.xlsx"
    workbook_path = Path(path).expanduser().resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Test data workbook not found: {workbook_path}")

    try:
        workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    except (InvalidFileException, OSError, zipfile.BadZipFile) as error:
        raise ValueError(f"Could not read test data workbook: {workbook_path}") from error

    sheet = workbook.active
    headers = tuple(sheet.cell(row=1, column=column).value for column in range(1, 5))
    if headers != EXPECTED_HEADERS:
        workbook.close()
        expected = ", ".join(EXPECTED_HEADERS)
        actual = ", ".join(str(header) for header in headers)
        raise ValueError(f"Expected Excel headers [{expected}], found [{actual}]")

    vegetables = []
    quantities = []
    try:
        for row_number in range(2, sheet.max_row + 1):
            name = sheet.cell(row=row_number, column=1).value
            quantity_value = sheet.cell(row=row_number, column=2).value
            if name is None or not str(name).strip():
                continue
            try:
                quantity = int(quantity_value)
                is_integer = float(quantity_value) == quantity
            except (TypeError, ValueError, OverflowError):
                is_integer = False
                quantity = 0
            if isinstance(quantity_value, bool) or not is_integer or quantity <= 0:
                raise ValueError(
                    f"Invalid quantity at row {row_number}: expected a positive integer"
                )
            vegetables.append(str(name).strip())
            quantities.append(quantity)

        if not vegetables:
            raise ValueError("Test data workbook contains no product rows")

        promocode = sheet.cell(row=2, column=3).value or ""
        country = sheet.cell(row=2, column=4).value or ""
    finally:
        workbook.close()

    return {
        "vegetables": vegetables,
        "quantities": quantities,
        "promocode": str(promocode).strip(),
        "country": str(country).strip()
    }